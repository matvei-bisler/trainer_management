"""
Проверка банка вопросов перед публикацией.

    python check_bank.py          — проверить Excel и показать проблемы
    python check_bank.py --fill   — ещё и заполнить пустые qid (Excel будет перезаписан)

Код выхода 1, если есть ошибки: такие вопросы не попадут в тренажёр и новые тесты.
"""

import sys
from collections import Counter

import openpyxl

from lib.bank import EXCEL_PATH, SHEET, _cell, load_bank, make_qid


def fill_qids(path=EXCEL_PATH) -> int:
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    header = [c.value for c in ws[1]]
    if "qid" in header:
        col = header.index("qid") + 1
    else:
        col = len([h for h in header if h is not None]) + 1
        ws.cell(row=1, column=col, value="qid")
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14
    names = {h: i + 1 for i, h in enumerate(header) if h}

    used = Counter(_cell(ws.cell(r, col).value) for r in range(2, ws.max_row + 1))
    filled = 0
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col)
        text = _cell(ws.cell(r, names["Q"]).value)
        if _cell(cell.value) or not text:
            continue
        # тот же хэш, что приложение подставляет для пустого qid: ссылки на вопрос не рвутся
        qid = make_qid(text, _cell(ws.cell(r, names["Q_opt"]).value))
        base, n = qid, 1
        while used[qid]:  # одинаковый текст — вторая копия получает суффикс
            qid, n = f"{base}-{n}", n + 1
        used[qid] += 1
        cell.value = qid
        filled += 1
    for r in range(1, ws.max_row + 1):
        ws.cell(r, col).number_format = "@"  # текст: Excel не превратит qid в число
    wb.save(path)
    return filled


def main() -> int:
    if "--fill" in sys.argv:
        print(f"Заполнено qid: {fill_qids()}")
    questions, issues = load_bank()
    print(f"Вопросов в банке: {len(questions)}")
    errors = [i for i in issues if i["level"] == "error"]
    warnings = [i for i in issues if i["level"] == "warning"]
    for title, items in (("ОШИБКИ", errors), ("Предупреждения", warnings)):
        if items:
            print(f"\n{title} ({len(items)}):")
            for i in items:
                print(f"  строка {i['row']:>4}  {i['qid'] or '':<12} {i['message']}")
    if not issues:
        print("Проблем не найдено.")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
